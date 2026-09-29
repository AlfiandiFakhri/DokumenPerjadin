import io
import os
import copy
import datetime
import zipfile
import streamlit as st
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.shared import Inches, Pt, RGBColor

# Konfigurasi halaman Streamlit
st.set_page_config(
    page_title="Nota Dinas & Surat Tugas Perjadin",
    page_icon="📋",
    layout="wide",
)

# ==============================================================================
# DATABASE PEGAWAI & JABATAN
# ==============================================================================
DATA_PEGAWAI_JABATAN = {
    "Hengki Andrianto": "Asisten Deputi Pembiayaan dan Investasi Usaha Menengah",
    "Ira Dunggio": "Kepala Bidang Renbang Pembiayaan dan Investasi Usaha Menengah",
    "Nadia Wulandari": "Perencana Ahli Muda",
    "Nurmina": "Pengembang Kewirausahaan Ahli Muda",
    "Feris Banda Mukti": "Analis Kebijakan Ahli Pertama",
    "Fakhri Alfiandi": "Pengembang Kewirausahaan Ahli Pertama",
    "Iqbal Mumtaz Fahmy": "Pengembang Kewirausahaan Ahli Pertama",
    "Tatag Suryo Pambudi": "Pengembang Kewirausahaan Ahli Pertama",
    "Maya Annisa": "Pengembang Kewirausahaan Ahli Pertama",
    "Beny Johani": "Pelaksana",
    "Lainnya (Ketik Manual...)": "Lainnya"
}

# ==============================================================================
# DATABASE PROVINSI & KABUPATEN/KOTA (38 PROVINSI)
# ==============================================================================
DATA_WILAYAH = {
    "Nanggroe Aceh Darussalam (NAD)": ["Banda Aceh", "Sabang", "Lhokseumawe", "Langsa", "Subulussalam", "Kab. Aceh Besar", "Kab. Pidie", "Kab. Bireuen", "Kab. Aceh Utara", "Kab. Aceh Timur", "Kab. Aceh Tengah", "Kab. Aceh Barat", "Kab. Aceh Selatan", "Kab. Aceh Singkil", "Kab. Aceh Tamiang", "Kab. Gayo Lues"],
    "Sumatera Utara": ["Medan", "Binjai", "Tebing Tinggi", "Pematangsiantar", "Tanjungbalai", "Sibolga", "Padangsidimpuan", "Gunungsitoli", "Kab. Deli Serdang", "Kab. Langkat", "Kab. Karo", "Kab. Simalungun", "Kab. Asahan", "Kab. Tapanuli Utara", "Kab. Tapanuli Selatan", "Kab. Tapanuli Tengah", "Kab. Nias"],
    "Sumatera Barat": ["Padang", "Bukittinggi", "Payakumbuh", "Solok", "Sawahlunto", "Padang Panjang", "Pariaman", "Kab. Agam", "Kab. Pesisir Selatan", "Kab. Lima Puluh Kota", "Kab. Tanah Datar", "Kab. Padang Pariaman", "Kab. Solok Selatan", "Kab. Dharmasraya"],
    "Riau": ["Pekanbaru", "Dumai", "Kab. Kampar", "Kab. Rokan Hulu", "Kab. Rokan Hilir", "Kab. Pelalawan", "Kab. Siak", "Kab. Bengkalis", "Kab. Indragiri Hilir", "Kab. Indragiri Hulu", "Kab. Kuantan Singingi"],
    "Kepulauan Riau": ["Batam", "Tanjungpinang", "Kab. Bintan", "Kab. Karimun", "Kab. Natuna", "Kab. Kepulauan Anambas", "Kab. Lingga"],
    "Jambi": ["Jambi", "Sungai Penuh", "Kab. Batanghari", "Kab. Bungo", "Kab. Kerinci", "Kab. Merangin", "Kab. Muaro Jambi", "Kab. Sarolangun", "Kab. Tanjung Jabung Barat", "Kab. Tanjung Jabung Timur", "Kab. Tebo"],
    "Bengkulu": ["Bengkulu", "Kab. Bengkulu Selatan", "Kab. Bengkulu Tengah", "Kab. Bengkulu Utara", "Kab. Kaur", "Kab. Kepahiang", "Kab. Lebong", "Kab. Mukomuko", "Kab. Rejang Lebong", "Kab. Seluma"],
    "Sumatera Selatan": ["Palembang", "Prabumulih", "Pagar Alam", "Lubuklinggau", "Kab. Banyuasin", "Kab. Empat Lawang", "Kab. Lahat", "Kab. Muara Enim", "Kab. Musi Banyuasin", "Kab. Musi Rawas", "Kab. Ogan Ilir", "Kab. Ogan Komering Ilir", "Kab. Ogan Komering Ulu"],
    "Kepulauan Bangka Belitung": ["Pangkalpinang", "Kab. Bangka", "Kab. Bangka Barat", "Kab. Bangka Selatan", "Kab. Bangka Tengah", "Kab. Belitung", "Kab. Belitung Timur"],
    "Lampung": ["Bandar Lampung", "Metro", "Kab. Lampung Barat", "Kab. Lampung Selatan", "Kab. Lampung Tengah", "Kab. Lampung Timur", "Kab. Lampung Utara", "Kab. Mesuji", "Kab. Pesawaran", "Kab. Pringsewu", "Kab. Tanggamus", "Kab. Tulang Bawang", "Kab. Way Kanan"],
    "Banten": ["Serang", "Cilegon", "Tangerang", "Tangerang Selatan", "Kab. Lebak", "Kab. Pandeglang", "Kab. Serang", "Kab. Tangerang"],
    "DKI Jakarta": ["Jakarta Pusat", "Jakarta Utara", "Jakarta Barat", "Jakarta Selatan", "Jakarta Timur", "Kepulauan Seribu"],
    "Jawa Barat": ["Bandung", "Bekasi", "Bogor", "Cimahi", "Cirebon", "Depok", "Sukabumi", "Tasikmalaya", "Banjar", "Kab. Bandung", "Kab. Bandung Barat", "Kab. Bekasi", "Kab. Bogor", "Kab. Ciamis", "Kab. Cianjur", "Kab. Cirebon", "Kab. Garut", "Kab. Indramayu", "Kab. Karawang", "Kab. Kuningan", "Kab. Majalengka", "Kab. Pangandaran", "Kab. Purwakarta", "Kab. Subang", "Kab. Sukabumi", "Kab. Sumedang", "Kab. Tasikmalaya"],
    "Jawa Tengah": ["Semarang", "Surakarta (Solo)", "Magelang", "Pekalongan", "Salatiga", "Tegal", "Kab. Banyumas", "Kab. Batang", "Kab. Blora", "Kab. Boyolali", "Kab. Brebes", "Kab. Cilacap", "Kab. Demak", "Kab. Grobogan", "Kab. Jepara", "Kab. Karanganyar", "Kab. Kebumen", "Kab. Kendal", "Kab. Klaten", "Kab. Kudus", "Kab. Magelang", "Kab. Pati", "Kab. Pekalongan", "Kab. Pemalang", "Kab. Purbalingga", "Kab. Purworejo", "Kab. Rembang", "Kab. Semarang", "Kab. Sragen", "Kab. Sukoharjo", "Kab. Tegal", "Kab. Temanggung", "Kab. Wonogiri", "Kab. Wonosobo"],
    "DI Yogyakarta": ["Yogyakarta", "Kab. Bantul", "Kab. Gunungkidul", "Kab. Kulon Progo", "Kab. Sleman"],
    "Jawa Timur": ["Surabaya", "Malang", "Batu", "Kediri", "Blitar", "Madiun", "Mojokerto", "Pasuruan", "Probolinggo", "Kab. Bangkalan", "Kab. Banyuwangi", "Kab. Blitar", "Kab. Bojonegoro", "Kab. Bondowoso", "Kab. Gresik", "Kab. Jember", "Kab. Jombang", "Kab. Kediri", "Kab. Lamongan", "Kab. Lumajang", "Kab. Madiun", "Kab. Magetan", "Kab. Malang", "Kab. Mojokerto", "Kab. Nganjuk", "Kab. Ngawi", "Kab. Pacitan", "Kab. Pamekasan", "Kab. Pasuruan", "Kab. Ponorogo", "Kab. Probolinggo", "Kab. Sampang", "Kab. Sidoarjo", "Kab. Situbondo", "Kab. Sumenep", "Kab. Trenggalek", "Kab. Tuban", "Kab. Tulungagung"],
    "Bali": ["Denpasar", "Kab. Badung", "Kab. Bangli", "Kab. Buleleng", "Kab. Gianyar", "Kab. Jembrana", "Kab. Karangasem", "Kab. Klungkung", "Kab. Tabanan"],
    "Nusa Tenggara Barat (NTB)": ["Mataram", "Bima", "Kab. Bima", "Kab. Dompu", "Kab. Lombok Barat", "Kab. Lombok Tengah", "Kab. Lombok Timur", "Kab. Lombok Utara", "Kab. Sumbawa", "Kab. Sumbawa Barat"],
    "Nusa Tenggara Timur (NTT)": ["Kupang", "Kab. Alor", "Kab. Belu", "Kab. Ende", "Kab. Flores Timur", "Kab. Kupang", "Kab. Lembata", "Kab. Malaka", "Kab. Manggarai", "Kab. Manggarai Barat", "Kab. Manggarai Timur", "Kab. Nagekeo", "Kab. Ngada", "Kab. Rote Ndao", "Kab. Sabu Raijua", "Kab. Sikka", "Kab. Sumba Barat", "Kab. Sumba Barat Daya", "Kab. Sumba Tengah", "Kab. Sumba Timur", "Kab. Timor Tengah Selatan", "Kab. Timor Tengah Utara"],
    "Kalimantan Barat": ["Pontianak", "Singkawang", "Kab. Bengkayang", "Kab. Kapuas Hulu", "Kab. Kayong Utara", "Kab. Ketapang", "Kab. Kubu Raya", "Kab. Landak", "Kab. Melawi", "Kab. Mempawah", "Kab. Sambas", "Kab. Sanggau", "Kab. Sekadau", "Kab. Sintang"],
    "Kalimantan Tengah": ["Palangka Raya", "Kab. Barito Selatan", "Kab. Barito Timur", "Kab. Barito Utara", "Kab. Gunung Mas", "Kab. Kapuas", "Kab. Katingan", "Kab. Kotawaringin Barat", "Kab. Kotawaringin Timur", "Kab. Lamandau", "Kab. Murung Raya", "Kab. Pulang Pisau", "Kab. Sukamara", "Kab. Seruyan"],
    "Kalimantan Selatan": ["Banjarmasin", "Banjarbaru", "Kab. Balangan", "Kab. Banjar", "Kab. Barito Kuala", "Kab. Hulu Sungai Selatan", "Kab. Hulu Sungai Tengah", "Kab. Hulu Sungai Utara", "Kab. Kotabaru", "Kab. Tabalong", "Kab. Tanah Bumbu", "Kab. Tanah Laut", "Kab. Tapin"],
    "Kalimantan Timur": ["Samarinda", "Balikpapan", "Bontang", "Nusantara (IKN)", "Kab. Berau", "Kab. Kutai Barat", "Kab. Kutai Kartanegara", "Kab. Kutai Timur", "Kab. Mahakam Ulu", "Kab. Paser", "Kab. Penajam Paser Utara"],
    "Kalimantan Utara": ["Tarakan", "Kab. Bulungan", "Kab. Malinau", "Kab. Nunukan", "Kab. Tana Tidung"],
    "Sulawesi Utara": ["Manado", "Bitung", "Tomohon", "Kotamobagu", "Kab. Bolaang Mongondow", "Kab. Minahasa", "Kab. Minahasa Selatan", "Kab. Minahasa Utara", "Kab. Kepulauan Sangihe", "Kab. Kepulauan Talaud"],
    "Gorontalo": ["Gorontalo", "Kab. Boalemo", "Kab. Bone Bolango", "Kab. Gorontalo", "Kab. Gorontalo Utara", "Kab. Pohuwato"],
    "Sulawesi Tengah": ["Palu", "Kab. Banggai", "Kab. Donggala", "Kab. Morowali", "Kab. Poso", "Kab. Sigi", "Kab. Tojo Una-Una", "Kab. Tolitoli", "Kab. Parigi Moutong"],
    "Sulawesi Barat": ["Kab. Majene", "Kab. Mamasa", "Kab. Mamuju", "Kab. Mamuju Tengah", "Kab. Pasangkayu", "Kab. Polewali Mandar"],
    "Sulawesi Selatan": ["Makassar", "Parepare", "Palopo", "Kab. Bantaeng", "Kab. Barru", "Kab. Bone", "Kab. Bulukumba", "Kab. Enrekang", "Kab. Gowa", "Kab. Jeneponto", "Kab. Luwu", "Kab. Luwu Timur", "Kab. Luwu Utara", "Kab. Maros", "Kab. Pangkajene dan Kepulauan", "Kab. Pinrang", "Kab. Selayar", "Kab. Sinjai", "Kab. Sidenreng Rappang", "Kab. Soppeng", "Kab. Takalar", "Kab. Tana Toraja", "Kab. Toraja Utara", "Kab. Wajo"],
    "Sulawesi Tenggara": ["Kendari", "Baubau", "Kab. Bombana", "Kab. Buton", "Kab. Kolaka", "Kab. Konawe", "Kab. Muna", "Kab. Wakatobi"],
    "Maluku": ["Ambon", "Tual", "Kab. Buru", "Kab. Kepulauan Aru", "Kab. Maluku Tengah", "Kab. Maluku Tenggara", "Kab. Seram Bagian Barat", "Kab. Seram Bagian Timur"],
    "Maluku Utara": ["Ternate", "Tidore Kepulauan", "Kab. Halmahera Barat", "Kab. Halmahera Selatan", "Kab. Halmahera Tengah", "Kab. Halmahera Timur", "Kab. Halmahera Utara", "Kab. Kepulauan Sula", "Kab. Pulau Morotai"],
    "Papua Barat Daya": ["Sorong", "Kab. Sorong", "Kab. Sorong Selatan", "Kab. Raja Ampat", "Kab. Tambrauw", "Kab. Maybrat"],
    "Papua Barat": ["Kab. Manokwari", "Kab. Fakfak", "Kab. Kaimana", "Kab. Teluk Bintuni", "Kab. Teluk Wondama", "Kab. Pegunungan Arfak", "Kab. Manokwari Selatan"],
    "Papua Tengah": ["Kab. Nabire", "Kab. Mimika", "Kab. Paniai", "Kab. Dogiyai", "Kab. Deiyai", "Kab. Intan Jaya", "Kab. Puncak", "Kab. Puncak Jaya"],
    "Papua": ["Jayapura", "Kab. Jayapura", "Kab. Keerom", "Kab. Sarmi", "Kab. Mamberamo Raya", "Kab. Biak Numfor", "Kab. Supiori", "Kab. Kepulauan Yapen", "Kab. Waropen"],
    "Papua Pegunungan": ["Kab. Jayawijaya", "Kab. Lanny Jaya", "Kab. Mamberamo Tengah", "Kab. Nduga", "Kab. Tolikara", "Kab. Yahukimo", "Kab. Yalimo", "Kab. Pegunungan Bintang"],
    "Papua Selatan": ["Kab. Merauke", "Kab. Boven Digoel", "Kab. Mappi", "Kab. Asmat"]
}

BULAN_INDO = {
    1: "Januari", 2: "Februari", 3: "Maret", 4: "April", 
    5: "Mei", 6: "Juni", 7: "Juli", 8: "Agustus", 
    9: "September", 10: "Oktober", 11: "November", 12: "Desember"
}

def set_document_font(doc, font_name="Arial", size_pt=12):
    style = doc.styles["Normal"]
    font = style.font
    font.name = font_name
    font.size = Pt(size_pt)
    font.color.rgb = RGBColor(0, 0, 0)

def setup_document_layout(section, kop_path, watermark_path, footer_path):
    section.header_distance = Inches(0.0)
    section.footer_distance = Inches(0.1)

    header = section.header
    hp_kop = header.paragraphs[0]
    hp_kop.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists(kop_path):
        try:
            hp_kop.add_run().add_picture(kop_path, width=Inches(3.2))
        except Exception as e:
            st.warning(f"Gagal memuat gambar kop surat: {e}")
    if os.path.exists(watermark_path):
        try:
            hp_wm = header.add_paragraph()
            run_wm = hp_wm.add_run()
            run_wm.add_picture(watermark_path, width=Inches(6.5)) 
            drawing = run_wm._r.xpath('.//w:drawing')[0]
            inline = drawing.xpath('.//wp:inline')[0]
            extent = inline.xpath('.//wp:extent')[0]
            cx = extent.get('cx')
            cy = extent.get('cy')
            docPr = inline.xpath('.//wp:docPr')[0]
            docPr_id = docPr.get('id')
            docPr_name = docPr.get('name')
            graphic = copy.deepcopy(inline.xpath('.//a:graphic')[0])
            anchor_xml = f'''
            <wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
                       xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
                       distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="0"
                       behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">
                <wp:simplePos x="0" y="0"/>
                <wp:positionH relativeFrom="page">
                    <wp:align>center</wp:align>
                </wp:positionH>
                <wp:positionV relativeFrom="page">
                    <wp:align>center</wp:align>
                </wp:positionV>
                <wp:extent cx="{cx}" cy="{cy}"/>
                <wp:effectExtent l="0" t="0" r="0" b="0"/>
                <wp:wrapNone/>
                <wp:docPr id="{docPr_id}" name="{docPr_name}"/>
                <wp:cNvGraphicFramePr>
                    <a:graphicFrameLocks noChangeAspect="1"/>
                </wp:cNvGraphicFramePr>
            </wp:anchor>
            '''
            anchor = parse_xml(anchor_xml)
            anchor.append(graphic)
            drawing.replace(inline, anchor)
        except Exception as e:
            st.warning(f"Gagal memuat watermark floating: {e}")
    footer = section.footer
    hp_footer = footer.paragraphs[0]
    hp_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists(footer_path):
        try:
            hp_footer.add_run().add_picture(footer_path, width=Inches(5.8))
        except Exception as e:
            st.warning(f"Gagal memuat gambar footer: {e}")

# ==============================================================================
# FUNGSI GENERATOR NOTA DINAS (A4, Arial 12)
# ==============================================================================
def generate_nota_dinas(data):
    doc = Document()
    set_document_font(doc, "Arial", 12)

    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        setup_document_layout(section, "kop_surat.png", "watermark.png", "footer_surat.png")

    p_judul = doc.add_paragraph()
    p_judul.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_judul.paragraph_format.space_after = Pt(2)
    r_judul = p_judul.add_run("NOTA DINAS")
    r_judul.bold = True
    r_judul.font.size = Pt(14)
    r_judul.font.name = "Arial"

    p_nomor = doc.add_paragraph()
    p_nomor.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_nomor.paragraph_format.space_after = Pt(6)
    r_nomor = p_nomor.add_run(f"Nomor: {data['nomor_memo']}")
    r_nomor.font.size = Pt(12)
    r_nomor.font.name = "Arial"

    header_table = doc.add_table(rows=3, cols=3)
    for cell in header_table.columns[0].cells:
        cell.width = Inches(1.1) 
    for cell in header_table.columns[1].cells:
        cell.width = Inches(0.2) 
    for cell in header_table.columns[2].cells:
        cell.width = Inches(4.9) 

    headers_data = [
        ("Kepada Yth.", ":", data['tujuan_memo']),
        ("Dari", ":", data['pengirim_memo']),
        ("Perihal", ":", "Permohonan Penerbitan ST"),
    ]

    for i, (label, separator, val) in enumerate(headers_data):
        cell_label = header_table.cell(i, 0)
        cell_separator = header_table.cell(i, 1)
        cell_val = header_table.cell(i, 2)
        
        p_label = cell_label.paragraphs[0]
        p_label.paragraph_format.space_after = Pt(0)
        p_label.paragraph_format.line_spacing = 1.0 
        p_label.add_run(label).font.name = "Arial"
        
        p_separator = cell_separator.paragraphs[0]
        p_separator.paragraph_format.space_after = Pt(0)
        p_separator.paragraph_format.line_spacing = 1.0
        p_separator.add_run(separator).font.name = "Arial"
        
        p_val = cell_val.paragraphs[0]
        p_val.paragraph_format.space_after = Pt(0)
        p_val.paragraph_format.line_spacing = 1.0
        p_val.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_val.add_run(val).font.name = "Arial"

    p_spasi = doc.add_paragraph()
    p_spasi.paragraph_format.space_after = Pt(2) 

    p_pembuka = doc.add_paragraph()
    p_pembuka.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_pembuka.paragraph_format.line_spacing = 1.50
    p_pembuka.paragraph_format.space_after = Pt(4)
    p_pembuka.paragraph_format.first_line_indent = Inches(0.5) 
    
    r_pembuka = p_pembuka.add_run(
        f"Dalam rangka {data['maksud_kegiatan']} dengan (MAK {data['mak_anggaran']}). Bersama ini "
        "disampaikan rencana perjalanan dinas pada Asisten Deputi Pembiayaan dan "
        "Investasi Usaha Menengah Tahun Anggaran 2026, dan dimohonkan dapat "
        "menerbitkan Surat Perintah Tugas atas nama sebagai berikut:"
    )
    r_pembuka.font.name = "Arial"

    table = doc.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    table.autofit = False 
    table.allow_autofit = False 
    table.alignment = WD_TABLE_ALIGNMENT.CENTER  
    
    col_widths = [
        Inches(0.40), # No
        Inches(1.82), # Nama
        Inches(0.44), # Gol
        Inches(0.71), # Waktu
        Inches(1.48), # Tujuan
        Inches(1.69), # Tanggal Dinas
        Inches(1.02)  # Kendaraan
    ]
    
    hdr_cells = table.rows[0].cells
    headers_title = ["No", "Nama", "Gol", "Waktu", "Tujuan", "Tanggal Dinas", "Kendaraan"]
    
    for i, title_text in enumerate(headers_title):
        hdr_cells[i].text = title_text
        hdr_cells[i].width = col_widths[i]
        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER 
        for paragraph in hdr_cells[i].paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(1) 
            paragraph.paragraph_format.space_after = Pt(1)  
            for run in paragraph.runs:
                run.bold = True
                run.font.name = "Arial"
                run.font.size = Pt(12)

    for idx, peg in enumerate(data['pegawai_list'], start=1):
        row_cells = table.add_row().cells
        row_values = [str(idx), peg['nama'], peg['gol'], f"{peg['waktu_hari']} Hari", peg['kab_final'], peg['tgl_str'], peg['kend_final']]
        for col_idx in range(7):
            row_cells[col_idx].text = row_values[col_idx]
            row_cells[col_idx].width = col_widths[col_idx]
            row_cells[col_idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            
            for paragraph in row_cells[col_idx].paragraphs:
                paragraph.paragraph_format.space_before = Pt(1)
                paragraph.paragraph_format.space_after = Pt(1)
                paragraph.paragraph_format.line_spacing = 1.0
                
                if col_idx in [0, 2, 3, 4]: 
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(12)

    p_penutup = doc.add_paragraph()
    p_penutup.paragraph_format.space_before = Pt(4)
    p_penutup.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_penutup.paragraph_format.line_spacing = 1.0
    p_penutup.paragraph_format.space_after = Pt(2)
    p_penutup.paragraph_format.first_line_indent = Inches(0.5) 
    r_penutup = p_penutup.add_run("\nDemikian disampaikan, atas kerjasamanya diucapkan terima kasih.")
    r_penutup.font.name = "Arial"

    ttd_table = doc.add_table(rows=1, cols=2)
    ttd_table.autofit = False 
    ttd_table.allow_autofit = False 
    for cell in ttd_table.columns[0].cells:
        cell.width = Inches(3.27) 
    for cell in ttd_table.columns[1].cells:
        cell.width = Inches(3.00) 
        
    cell_kiri = ttd_table.cell(0, 0)
    cell_kanan = ttd_table.cell(0, 1)
    
    p_kiri = cell_kiri.paragraphs[0]
    p_kiri.paragraph_format.line_spacing = 1.0
    r_kiri = p_kiri.add_run("\n\n\n\n\n\nKabid 3.1.1.........................")
    r_kiri.font.name = "Arial"
    r_kiri.font.size = Pt(12)
    
    p_kanan = cell_kanan.paragraphs[0]
    p_kanan.paragraph_format.line_spacing = 1.0
    p_kanan.alignment = WD_ALIGN_PARAGRAPH.LEFT 
    
    r_tgl = p_kanan.add_run(f"\nJakarta, {data['tgl_hari']} {data['tgl_bulan']} {data['tgl_tahun']}\n")
    r_tgl.font.name = "Arial"
    r_tgl.font.size = Pt(12)
    
    r_jab = p_kanan.add_run("Asisten Deputi Pembiayaan dan\nInvestasi Usaha Menengah\n\n")
    r_jab.font.name = "Arial"
    r_jab.font.size = Pt(12)
    
    r_nama = p_kanan.add_run("\n\n\n\nHengki Andrianto\n")
    r_nama.bold = True
    r_nama.font.name = "Arial"
    r_nama.font.size = Pt(12)
    
    r_nip = p_kanan.add_run("NIP. 19820512 200912 1 001")
    r_nip.font.name = "Arial"
    r_nip.font.size = Pt(12)

    p_tembusan = doc.add_paragraph()
    p_tembusan.paragraph_format.space_before = Pt(4) 
    p_tembusan.paragraph_format.line_spacing = 1.0   
    r_tembusan = p_tembusan.add_run(
        "\n\nTembusan Yth:\n"
        "1. Deputi Bidang Usaha Menengah;\n"
        "2. PPK Satuan Kerja Deputi Bidang Usaha Menengah;\n"
        "3. PPSPM Satuan Kerja Deputi Bidang Usaha Menengah;\n"
        "4. BP Satuan Kerja Deputi Bidang Usaha Menengah."
    )
    r_tembusan.font.name = "Arial"
    r_tembusan.font.size = Pt(12) 

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


# ==============================================================================
# FUNGSI GENERATOR SURAT TUGAS (F4, Bookman Old Style 12 pt)
# MURNI TANPA TABEL (MENGGUNAKAN TABULASI & INDENTASI PARAGRAPH)
# ==============================================================================
def generate_surat_tugas(data):
    doc = Document()
    set_document_font(doc, "Bookman Old Style", 12)

    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(13.0)
        section.top_margin = Inches(0.5)     
        section.bottom_margin = Inches(0.4)  
        section.left_margin = Inches(0.6)    
        section.right_margin = Inches(0.3)   
        setup_document_layout(section, "kop_surat.png", "watermark.png", "footer_surat.png")

    # Judul Surat Tugas 
    p_judul = doc.add_paragraph()
    p_judul.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_judul.paragraph_format.space_after = Pt(1)
    p_judul.paragraph_format.line_spacing = 1.5
    r_judul = p_judul.add_run("SURAT TUGAS")
    r_judul.bold = True
    r_judul.font.size = Pt(14)
    r_judul.font.name = "Bookman Old Style"

    p_nomor = doc.add_paragraph()
    p_nomor.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_nomor.paragraph_format.space_after = Pt(6)
    p_nomor.paragraph_format.line_spacing = 1.5
    r_nomor = p_nomor.add_run(f"NOMOR: {data['nomor_st']}")
    r_nomor.font.size = Pt(12)
    r_nomor.font.name = "Bookman Old Style"

    # Posisi Tabulasi Standar untuk Bagian Pegawai & Poin III s.d. V
    TAB_POS = Inches(2.50)  
    
    # I. DIPERINTAHKAN KEPADA 
    p_diperintah = doc.add_paragraph()
    p_diperintah.paragraph_format.space_after = Pt(2)
    p_diperintah.paragraph_format.line_spacing = 1.5
    r_diperintah = p_diperintah.add_run("I. DIPERINTAHKAN KEPADA")
    r_diperintah.bold = False
    r_diperintah.font.name = "Bookman Old Style"
    r_diperintah.font.size = Pt(12)

    # Padding list pegawai agar selalu mencantumkan nomor 1 sampai 4
    pegawai_list_padded = list(data['pegawai_list'])
    while len(pegawai_list_padded) < 4:
        pegawai_list_padded.append({"nama": "", "jabatan": ""})

    INDENT_ANGKA = Inches(0.3)
    for idx, peg in enumerate(pegawai_list_padded[:4], start=1):
        # Baris Nama
        p_nama = doc.add_paragraph()
        p_nama.paragraph_format.left_indent = TAB_POS
        p_nama.paragraph_format.first_line_indent = -TAB_POS + INDENT_ANGKA
        p_nama.paragraph_format.tab_stops.add_tab_stop(TAB_POS)
        p_nama.paragraph_format.space_after = Pt(1)
        p_nama.paragraph_format.line_spacing = 1.5
        
        r_nama = p_nama.add_run(f"{idx}.  Nama\t: {peg['nama']}")
        r_nama.font.name = "Bookman Old Style"
        r_nama.font.size = Pt(12)

        # Baris Jabatan
        p_jab = doc.add_paragraph()
        p_jab.paragraph_format.left_indent = TAB_POS
        p_jab.paragraph_format.first_line_indent = -TAB_POS + INDENT_ANGKA
        p_jab.paragraph_format.tab_stops.add_tab_stop(TAB_POS)
        p_jab.paragraph_format.space_after = Pt(3)
        p_jab.paragraph_format.line_spacing = 1.5
        
        r_jab = p_jab.add_run(f"     Jabatan\t: {peg['jabatan']}")
        r_jab.font.name = "Bookman Old Style"
        r_jab.font.size = Pt(12)

    # ==========================================================================
    # KHUSUS II. MAKSUD PERJALANAN (Menggunakan Tabulasi Ganda agar Wrap Rapi)
    # ==========================================================================
    TAB_COLON = Inches(2.50)  # Posisi titik dua (:)
    TAB_TEXT = Inches(2.60)  # Posisi awal teks (lurus di bawah kata "Koordinasi")

    p_ii = doc.add_paragraph()
    p_ii.paragraph_format.left_indent = TAB_TEXT
    p_ii.paragraph_format.first_line_indent = -TAB_TEXT
    p_ii.paragraph_format.tab_stops.add_tab_stop(TAB_COLON)
    p_ii.paragraph_format.tab_stops.add_tab_stop(TAB_TEXT)
    p_ii.paragraph_format.space_after = Pt(2)
    p_ii.paragraph_format.line_spacing = 1.5
    
    r_ii = p_ii.add_run(f"II. MAKSUD PERJALANAN\t:\t{data['maksud_kegiatan']}")
    r_ii.font.name = "Bookman Old Style"
    r_ii.font.size = Pt(12)
    r_ii.bold = False

    # ==========================================================================
    # III, IV & V (Menggunakan Format Standar TAB_POS Semula)
    # ==========================================================================
    other_details = [
        ("III. TUJUAN", f": {data['wilayah_tujuan_st']}"),
        ("IV. JANGKA WAKTU", f": {data['jangka_waktu_st']} Hari"),
        ("     Tanggal Berangkat", f": {data['tgl_berangkat_st']}"),
        ("     Tanggal Kembali", f": {data['tgl_kembali_st']}"),
        ("V. KETERANGAN", ": 1. Setelah selesai melakukan perjalanan dinas ini segera\n      membuat laporan kepada kami.\n  2. SPT ini berlaku sejak tanggal dikeluarkan.")
    ]

    for label, val in other_details:
        p_det = doc.add_paragraph()
        p_det.paragraph_format.left_indent = TAB_POS
        p_det.paragraph_format.first_line_indent = -TAB_POS
        p_det.paragraph_format.tab_stops.add_tab_stop(TAB_POS)
        p_det.paragraph_format.space_after = Pt(2)
        p_det.paragraph_format.line_spacing = 1.5
        
        r_det = p_det.add_run(f"{label}\t{val}")
        r_det.font.name = "Bookman Old Style"
        r_det.font.size = Pt(12)
        r_det.bold = False

    # ==========================================================================
    # TANDA TANGAN DINAMIS BERDASARKAN TUJUAN NOTA DINAS
    # ==========================================================================
    if "Deputi Bidang Usaha Menengah" in data['tujuan_memo'] and "Sekretaris" not in data['tujuan_memo']:
        jabatan_penandatangan = "Deputi Bidang Usaha Menengah,"
        nama_penandatangan = "Bagus Rachman"
        nip_penandatangan = "NIP. 19710501 199903 1 001"
    else:
        jabatan_penandatangan = "Sekretaris Deputi Bidang Usaha Menengah,"
        nama_penandatangan = "Christina Agustin"
        nip_penandatangan = "NIP. 19720801 199803 2 001"

    sig_indent = Inches(4.2)  # Menggeser blok tanda tangan ke sisi kanan halaman
    tab_sig = Inches(5.5)     # Posisi titik dua pada baris tanggal

    p_tgl1 = doc.add_paragraph()
    p_tgl1.paragraph_format.left_indent = sig_indent
    p_tgl1.paragraph_format.tab_stops.add_tab_stop(tab_sig)
    p_tgl1.paragraph_format.space_before = Pt(10)
    p_tgl1.paragraph_format.space_after = Pt(1)
    p_tgl1.paragraph_format.line_spacing = 1.0
    r = p_tgl1.add_run("Dikeluarkan di\t: Jakarta")
    r.font.name = "Bookman Old Style"
    r.font.size = Pt(12)

    p_tgl2 = doc.add_paragraph()
    p_tgl2.paragraph_format.left_indent = sig_indent
    p_tgl2.paragraph_format.tab_stops.add_tab_stop(tab_sig)
    p_tgl2.paragraph_format.space_after = Pt(14) 
    p_tgl2.paragraph_format.line_spacing = 1.0
    r = p_tgl2.add_run(f"Pada Tanggal\t:            {data['st_bulan']} {data['st_tahun']}")
    r.font.name = "Bookman Old Style"
    r.font.size = Pt(12)

    p_jab_ttd = doc.add_paragraph()
    p_jab_ttd.paragraph_format.left_indent = sig_indent
    p_jab_ttd.paragraph_format.space_after = Pt(36) 
    p_jab_ttd.paragraph_format.line_spacing = 1.0
    r = p_jab_ttd.add_run(jabatan_penandatangan)
    r.font.name = "Bookman Old Style"
    r.font.size = Pt(12)

    p_nama = doc.add_paragraph()
    p_nama.paragraph_format.left_indent = sig_indent
    p_nama.paragraph_format.space_after = Pt(0)
    p_nama.paragraph_format.line_spacing = 1.0
    r = p_nama.add_run(f"\n\n{nama_penandatangan}")
    r.bold = True
    r.font.name = "Bookman Old Style"
    r.font.size = Pt(12)

    p_nip = doc.add_paragraph()
    p_nip.paragraph_format.left_indent = sig_indent
    p_nip.paragraph_format.space_after = Pt(0)
    p_nip.paragraph_format.line_spacing = 1.0
    r = p_nip.add_run(nip_penandatangan)
    r.font.name = "Bookman Old Style"
    r.font.size = Pt(12)

    # Tembusan Surat Tugas
    p_temb_title = doc.add_paragraph()
    p_temb_title.paragraph_format.space_before = Pt(8)
    p_temb_title.paragraph_format.space_after = Pt(1)
    p_temb_title.paragraph_format.line_spacing = 1.0
    r = p_temb_title.add_run("\nTembusan:")
    r.font.name = "Bookman Old Style"
    r.font.size = Pt(12)

    tembusan_list = [
        "1. Pejabat Pembuat Komitmen;",
        "2. Pejabat Penguji dan Penandatangan SPM;",
        "3. Bendahara Pengeluaran; dan",
        "4. Arsip."
    ]
    for item in tembusan_list:
        p_item = doc.add_paragraph()
        p_item.paragraph_format.space_after = Pt(0)
        p_item.paragraph_format.line_spacing = 1.0
        r = p_item.add_run(item)
        r.font.name = "Bookman Old Style"
        r.font.size = Pt(12)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


# ==============================================================================
# TAMPILAN ANTARMUKA STREAMLIT
# ==============================================================================
st.title("Nota Dinas & Surat Tugas Perjadin")

col_setup1, col_setup2, col_setup3 = st.columns(3)
with col_setup1:
    jumlah_pegawai = st.number_input(
        "👥 Jumlah pegawai Pelaksana Dinas", 
        min_value=1, 
        max_value=100, 
        value=1, 
        step=1
    )

with col_setup2:
    selected_mak_full = st.selectbox(
        "💰 Pilih MAK Anggaran",
        [
            "7475.QDG.001.051.A.524111 (Koordinasi, Identifikasi, Persiapan Pelaksanaan Program Perluasan Pembiayaan KUR Klaster)",
            "7475.QDG.001.052.A.524111 (Koordinasi, Identifikasi, Persiapan Pelaksanaan Program Akselerasi Pembiayaan Inovatif bagi Usaha Menengah)"
        ]
    )

with col_setup3:
    jenis_output = st.multiselect(
        "📄 Pilih Dokumen yang Ingin Dibuat",
        ["Nota Dinas", "Surat Tugas (ST)"],
        default=["Nota Dinas", "Surat Tugas (ST)"]
    )

mak_anggaran = selected_mak_full.split(" (")[0]
maksud_kegiatan = selected_mak_full.split(" (")[1].replace(")", "")

DAFTAR_KENDARAAN = [
    "Pesawat",
    "Mobil",
    "BUS",
    "Kereta Api",
    "Travel",
    "Lainya (Bisa ketik manual)"
]

st.subheader("📌 1. Informasi Nomor Surat")
col1, col2, col3 = st.columns(3)
with col1:
    nomor_memo = st.text_input(
        "Nomor Nota Dinas",
        value="      /D.3.1.UMKM/TU.     /2026",
        placeholder="Contoh: 0000/D.3.1.UMKM/TU.00/2026",
    )
with col2:
    nomor_st = st.text_input(
        "Nomor Surat Tugas (ST)",
        value="    /ST/  D.3.UMKM/TU.    /2026",
        placeholder="Contoh:     /ST/  D.3.UMKM/TU.    /2026",
    )
with col3:
    tujuan_memo = st.selectbox(
        "Kepada Yth.",
        [
            "Deputi Bidang Usaha Menengah selaku Pejabat Penandatangan Surat Tugas (ST)",
            "Sekretaris Deputi Bidang Usaha Menengah selaku Pejabat Penandatangan Surat Tugas (ST)",
        ],
    )

pengirim_memo = "Asisten Deputi Pembiayaan dan Investasi Usaha Menengah"

st.markdown("---")
st.subheader("👥 2. Data Pegawai Pelaksana Perjalanan Dinas")

pegawai_data = []
for i in range(jumlah_pegawai):
    st.markdown(f"**👤 Pegawai {i+1}**")
    
    col_p1, col_p2, col_p3 = st.columns(3)
    with col_p1:
        pil_nama = st.selectbox(
            f"Pilih Nama Pegawai {i+1}", 
            list(DATA_PEGAWAI_JABATAN.keys()), 
            key=f"pil_nama_{i}"
        )
        
        if pil_nama == "Lainnya (Ketik Manual...)":
            nama_manual = st.text_input(f"Ketik Nama Pegawai {i+1}", placeholder="Nama Lengkap...", key=f"nama_man_{i}")
            jabatan_manual = st.text_input(f"Ketik Jabatan Pegawai {i+1}", placeholder="Jabatan Resmi...", key=f"jab_man_{i}")
        else:
            nama_manual = pil_nama
            jabatan_manual = DATA_PEGAWAI_JABATAN[pil_nama]

        gol = st.selectbox(f"Golongan {i+1}", ["III", "IV", "II", "-"], key=f"gol_{i}")
    
    with col_p2:
        prov = st.selectbox(f"Provinsi Tujuan {i+1}", list(DATA_WILAYAH.keys()), key=f"prov_{i}")
        
        opsi_kab = DATA_WILAYAH[prov] + ["Lainnya (Ketik Manual...)"]
        pil_kab = st.selectbox(f"Kabupaten/Kota Tujuan {i+1}", opsi_kab, key=f"pil_kab_{i}")
        
        if pil_kab == "Lainnya (Ketik Manual...)":
            kab_manual = st.text_input(f"Ketik Kabupaten/Kota {i+1}", placeholder="Nama Kab/Kota...", key=f"kab_man_{i}")
            kab_final = kab_manual
        else:
            kab_final = pil_kab
            
        today = datetime.date.today()
        tgl_input = st.date_input(
            f"Rentang Tanggal Dinas {i+1}",
            value=(today, today + datetime.timedelta(days=1)),
            key=f"tgl_{i}"
        )
        
        if isinstance(tgl_input, tuple):
            if len(tgl_input) == 2:
                d1, d2 = tgl_input
                jumlah_hari = (d2 - d1).days + 1
                if d1.month == d2.month and d1.year == d2.year:
                    tgl_str = f"{d1.day} s.d. {d2.day} {BULAN_INDO[d1.month]} {d1.year}"
                else:
                    tgl_str = f"{d1.day} {BULAN_INDO[d1.month]} s.d. {d2.day} {BULAN_INDO[d2.month]} {d2.year}"
                tgl_berangkat_str = f"{d1.day} {BULAN_INDO[d1.month]} {d1.year}"
                tgl_kembali_str = f"{d2.day} {BULAN_INDO[d2.month]} {d2.year}"
            elif len(tgl_input) == 1:
                d1 = tgl_input[0]
                jumlah_hari = 1
                tgl_str = f"{d1.day} {BULAN_INDO[d1.month]} {d1.year}"
                tgl_berangkat_str = tgl_str
                tgl_kembali_str = tgl_str
            else:
                jumlah_hari = 1
                tgl_str = ""
                tgl_berangkat_str = ""
                tgl_kembali_str = ""
        else:
            d1 = tgl_input
            jumlah_hari = 1
            tgl_str = f"{d1.day} {BULAN_INDO[d1.month]} {d1.year}"
            tgl_berangkat_str = tgl_str
            tgl_kembali_str = tgl_str

    with col_p3:
        st.markdown(f"**Waktu (Hari) {i+1}**")
        st.info(f"⏱️ **{jumlah_hari} Hari**")
        
        pil_kendaraan = st.selectbox(f"Kendaraan {i+1}", DAFTAR_KENDARAAN, key=f"pil_kend_{i}")
        
        if pil_kendaraan == "Lainya (Bisa ketik manual)":
            kend_manual = st.text_input(f"Ketik Kendaraan {i+1}", placeholder="Nama Kendaraan...", key=f"kend_man_{i}")
            kend_final = kend_manual
        else:
            kend_final = pil_kendaraan
    
    pegawai_data.append({
        "nama": nama_manual,
        "jabatan": jabatan_manual,
        "gol": gol,
        "waktu_hari": jumlah_hari,
        "prov": prov,
        "kab_final": kab_final,
        "tgl_str": tgl_str,
        "tgl_berangkat": tgl_berangkat_str,
        "tgl_kembali": tgl_kembali_str,
        "kend_final": kend_final
    })
    
    if i < jumlah_pegawai - 1:
        st.divider()

st.markdown("---")
st.subheader("📅 3. Tanggal Surat")

col_dt1, col_dt2, col_dt3 = st.columns(3)
with col_dt1:
    tgl_surat_input = st.date_input("Tanggal Nota Dinas", value=datetime.date.today(), key="tgl_surat_dinas")
    tgl_hari = tgl_surat_input.day
    tgl_bulan = BULAN_INDO[tgl_surat_input.month]
    tgl_tahun = tgl_surat_input.year

with col_dt2:
    st_bulan_pilih = st.selectbox(
        "Bulan Surat Tugas", 
        list(BULAN_INDO.values()), 
        index=datetime.date.today().month - 1, 
        key="st_bulan_pilih"
    )

with col_dt3:
    st_tahun_pilih = st.number_input(
        "Tahun Surat Tugas", 
        min_value=2024, 
        max_value=2030, 
        value=datetime.date.today().year, 
        key="st_tahun_pilih"
    )

# Tombol Eksekusi
submitted = st.button("🚀 PROSES DOKUMEN")

# --- PROSES PEMBUATAN DOKUMEN WORD ---
if submitted:
    if not jenis_output:
        st.error("⚠️ Mohon pilih minimal satu jenis dokumen yang ingin dibuat pada opsi di atas.")
    else:
        pegawai_list = []
        first_name = "Master_Baku"
        
        for idx, peg in enumerate(pegawai_data):
            if peg["nama"].strip(): 
                pegawai_list.append(peg)
                if idx == 0:
                    first_name = peg["nama"].strip().replace(" ", "_")

        wilayah_st = pegawai_data[0]["kab_final"] if len(pegawai_data) == 1 else pegawai_data[0]["prov"]

        data = {
            "nomor_memo": nomor_memo,
            "nomor_st": nomor_st,
            "tujuan_memo": tujuan_memo,
            "pengirim_memo": pengirim_memo,
            "maksud_kegiatan": maksud_kegiatan,
            "mak_anggaran": mak_anggaran,
            "pegawai_list": pegawai_list,
            "wilayah_tujuan_st": wilayah_st,
            "jangka_waktu_st": pegawai_data[0]["waktu_hari"],
            "tgl_berangkat_st": pegawai_data[0]["tgl_berangkat"],
            "tgl_kembali_st": pegawai_data[0]["tgl_kembali"],
            "tgl_hari": tgl_hari,
            "tgl_bulan": tgl_bulan,
            "tgl_tahun": tgl_tahun,
            "st_bulan": st_bulan_pilih,
            "st_tahun": st_tahun_pilih
        }

        st.success("🎉 Dokumen berhasil dibuat!")
        
        # Definisikan variabel nama kota dan tanggal untuk penamaan file
        nama_kota_clean = pegawai_data[0]["kab_final"].replace(" ", "_").replace("/", "-")
        rentang_tgl_clean = f"{pegawai_data[0]['tgl_berangkat']}_s.d._{pegawai_data[0]['tgl_kembali']}".replace(" ", "_")

        files_to_download = {}
        if "Nota Dinas" in jenis_output:
            buffer_nd = generate_nota_dinas(data)
            files_to_download[f"Nota_Dinas_Perjadin_{nama_kota_clean}_{rentang_tgl_clean}.docx"] = buffer_nd.getvalue()
                
        if "Surat Tugas (ST)" in jenis_output:
            buffer_st = generate_surat_tugas(data)
            files_to_download[f"Surat_Tugas_Perjadin_{nama_kota_clean}_{rentang_tgl_clean}.docx"] = buffer_st.getvalue()

        zip_filename = f"ND_dan_SPT_{nama_kota_clean}_{rentang_tgl_clean}.zip"

        zip_buffer = io.BytesIO()
        with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
            for filename, file_bytes in files_to_download.items():
                zip_file.writestr(filename, file_bytes)
        zip_buffer.seek(0)
        
        st.download_button(
            label="📥 Unduh Dokumen (.zip)",
            data=zip_buffer,
            file_name=zip_filename,
            mime="application/zip",
        )